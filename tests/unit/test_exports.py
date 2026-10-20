# -*- coding: utf-8 -*-
"""
tests/unit/test_exports.py — Exports XLSX/PDF du reporting (spec §15.3).

Vérifie que l'export reproduit EXACTEMENT le contenu Markdown (source de
vérité unique) sans masquer les statuts honnêtes, et que les fichiers sont
valides (XLSX relisible, PDF commençant par %PDF).
"""

import os
import pytest

from structcalc.reporting.exports import (
    parse_markdown,
    export_xlsx,
    export_pdf,
    export_project_summary_xlsx,
    export_project_summary_pdf,
)


MD = """# Document
## Verifications
| Check | Statut |
|---|---|
| compression_central | [?] A valider |
| flexion_composed | [X] Non conforme |
| rigide | [!] Verification non implemente |
## Autre
Texte avec **gras** et <tag> & ampersand.
"""


def test_parse_markdown_blocks():
    blocks = parse_markdown(MD)
    kinds = [b.kind for b in blocks]
    assert kinds == ["heading", "heading", "table", "heading", "text"]
    # la table : en-tête + 3 lignes de données (la séparation |---| est retirée)
    tbl = next(b for b in blocks if b.kind == "table")
    assert len(tbl.rows) == 4
    assert tbl.rows[0] == ["Check", "Statut"]
    assert tbl.rows[1][1] == "[?] A valider"


def test_export_xlsx_valid_and_faithful(tmp_path):
    p = str(tmp_path / "out.xlsx")
    result = export_xlsx(MD, p)
    assert result == p and os.path.exists(p) and os.path.getsize(p) > 0

    from openpyxl import load_workbook
    wb = load_workbook(p)
    # une feuille par section ##
    assert "Verifications" in wb.sheetnames and "Autre" in wb.sheetnames
    # les statuts honnêtes sont préservés à l'identique (source unique)
    vals = [c.value for row in wb["Verifications"].iter_rows() for c in row]
    assert "[?] A valider" in vals
    assert "[X] Non conforme" in vals
    assert "[!] Verification non implemente" in vals


def test_export_pdf_valid_header(tmp_path):
    p = str(tmp_path / "out.pdf")
    result = export_pdf(MD, p, title="Document")
    assert result == p and os.path.exists(p) and os.path.getsize(p) > 0
    with open(p, "rb") as f:
        assert f.read(5) == b"%PDF-"


def test_export_escapes_xml_specials(tmp_path):
    """Les caractères < > & ne doivent pas casser le rendu PDF."""
    md = "# T\nA < b > c & d\n"
    p = str(tmp_path / "esc.pdf")
    export_pdf(md, p)
    assert os.path.getsize(p) > 0


def test_project_summary_exports(tmp_path):
    class FakeProj:
        name = "Demo"
        def counts(self):
            return {"nodes": 4, "beams": 2}
        def data_gaps(self):
            return {"SC-BEAM-000001": ["section"]}

    xp = str(tmp_path / "proj.xlsx")
    pp = str(tmp_path / "proj.pdf")
    export_project_summary_xlsx(FakeProj(), xp)
    export_project_summary_pdf(FakeProj(), pp)
    assert os.path.getsize(xp) > 0 and os.path.getsize(pp) > 0
    # les données manquantes sont AFFICHÉES, jamais masquées (§15.3)
    from openpyxl import load_workbook
    wb = load_workbook(xp)
    vals = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row
            if c.value is not None]
    assert any("section" in v for v in vals)
